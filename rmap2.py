import subprocess
import sys

try:
    import openpyxl
except ModuleNotFoundError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])

import streamlit as st
import pandas as pd
import re
import io
import zipfile
from datetime import datetime
from openpyxl.styles import PatternFill, Font, Alignment

st.set_page_config(page_title="Renewal Mails Mapper", page_icon="🔖", layout="wide")

if 'mapping_results' not in st.session_state:
    st.session_state.mapping_results = None
if 'file_status' not in st.session_state:
    st.session_state.file_status = None

st.title("Renewal Mapping Dashboard")
st.markdown("Upload your Main File(s) and a Mapping Reference File to tag rows")

# =====================================================================
# === CORE MATCHING LOGIC (unchanged from the batch script) =========
# =====================================================================

def clean_text(text):
    text = str(text).lower()
    text = re.sub(r"\*+", " ", text)
    text = re.sub(r"[^a-z0-9 ]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def extract_parts(pattern):
    parts = re.split(r"\{\{.*?\}\}", str(pattern))
    parts = [p.strip() for p in parts if p.strip() != ""]

    if len(parts) == 1:
        return parts[0], ""
    elif len(parts) >= 2:
        return parts[0], parts[-1]
    else:
        return "", ""


# --- FILE UPLOAD SECTION ---
col1, col2 = st.columns(2)

with col1:
    st.subheader("1. Load Main File(s)")
    st.info("Upload (.csv, .xlsx) - all with same header structure")
    main_files = st.file_uploader(
        "Upload Main Sheet(s)",
        type=["csv", "xlsx"],
        key="main",
        accept_multiple_files=True
    )

    if main_files and len(main_files) > 13:
        st.error(f"❌ You uploaded {len(main_files)} files. Maximum allowed is 13 files. Please remove {len(main_files) - 13} file(s).")
        main_files = main_files[:13]
        st.info("Trimmed to first 13 files.")

with col2:
    st.subheader("2. Load Mapping Reference File")
    st.info("Upload mapping file (will be used for all main files)")
    mapping_file = st.file_uploader("Upload Mapping Patterns (.csv, .xlsx)", type=["csv", "xlsx"], key="map")

if main_files and mapping_file:

    @st.cache_data(ttl=3600)
    def load_uploaded_data(file_obj):
        if file_obj.name.endswith(".csv"):
            return pd.read_csv(file_obj)
        return pd.read_excel(file_obj)

    df_map = load_uploaded_data(mapping_file).copy()

    st.success(f"Successfully loaded {len(main_files)} main file(s) and mapping reference file!")
    st.info(f"Mapping file size: {len(df_map):,} rows")

    with st.expander("View Uploaded Files"):
        for idx, file in enumerate(main_files, 1):
            st.write(f"{idx}. **{file.name}**")

    # --- COLUMN MAPPER SELECTION ---
    st.subheader("Map Your Column Fields")

    col_sel1, col_sel2 = st.columns(2)

    with col_sel1:
        st.info("Identify the text/value column in your Main file(s):")
        first_file_df = load_uploaded_data(main_files[0])
        subject_col_name = st.selectbox(
            "Select Subject Column (Main File)",
            options=first_file_df.columns,
            key="subject_select"
        )

    with col_sel2:
        st.info("Identify the pattern column (with {{placeholders}}) from your Map file:")
        pattern_col = st.selectbox("Pattern Column (Map File)", options=df_map.columns, index=0 if len(df_map.columns) > 0 else 0)

    # --- DYNAMIC OUTPUT COLUMN SELECTION ---
    st.subheader("Select Output Mapping Columns")

    available_output_cols = [col for col in df_map.columns if col != pattern_col]

    if len(available_output_cols) > 0:
        selected_output_cols = st.multiselect(
            "Select which columns to map as output (you can select multiple)",
            options=available_output_cols,
            default=available_output_cols[:min(2, len(available_output_cols))]
        )
    else:
        st.warning("⚠️ No additional columns available in your mapping file for output!")
        selected_output_cols = []

    # --- CUSTOMIZABLE OUTPUT COLUMN NAMES ---
    if selected_output_cols:
        st.subheader("Customize Output Column Names")

        custom_output_names = {}
        cols_per_row = 2
        cols = st.columns(cols_per_row)

        for idx, col in enumerate(selected_output_cols):
            with cols[idx % cols_per_row]:
                custom_name = st.text_input(
                    f"Custom name for '{col}'",
                    value=f"Mapped_{col}",
                    key=f"custom_{col}"
                )
                custom_output_names[col] = custom_name

        # --- PROCESSING TRIGGER ---
        if st.button("Run Mapping Engine on All Files", type="primary"):

            all_results = []
            file_status_info = []

            overall_progress_bar = st.progress(0)
            overall_status_text = st.empty()

            for file_idx, main_file in enumerate(main_files, 1):
                file_name = main_file.name
                overall_status_text.write(f"📄 **Processing File {file_idx}/{len(main_files)}: {file_name}**")

                with st.status(f"Processing {file_name}...", expanded=False) as status:

                    try:
                        df_main = load_uploaded_data(main_file).copy()
                        df_main_original = df_main.copy()
                        status.write(f"✓ File loaded: {len(df_main):,} rows")

                        if subject_col_name not in df_main.columns:
                            st.error(f"❌ Column '{subject_col_name}' not found in {file_name}. Skipping this file.")
                            file_status_info.append({
                                "file": file_name,
                                "status": "❌ Failed",
                                "rows": 0,
                                "reason": f"Column '{subject_col_name}' not found"
                            })
                            continue

                        # === Step 1: extract before/after around {{placeholders}} and clean ===
                        status.write("Step 1/3: Extracting before/after pattern text and cleaning...")

                        df_map_temp = df_map.copy()
                        df_map_temp["before"], df_map_temp["after"] = zip(*df_map_temp[pattern_col].apply(extract_parts))
                        df_map_temp["before"] = df_map_temp["before"].apply(clean_text)
                        df_map_temp["after"] = df_map_temp["after"].apply(clean_text)

                        df_main["subject_clean"] = df_main[subject_col_name].apply(clean_text)

                        # Initialize output columns to "No Match" (same default as the batch script)
                        for col in selected_output_cols:
                            df_main[f"mapped_{col}"] = "No Match"
                        matched_flag = pd.Series(False, index=df_main.index)

                        # === Step 2: matching (substring, first-match-wins) ===
                        status.write("Step 2/3: Matching subjects against patterns...")
                        step_progress = st.progress(0.0)

                        total_patterns = len(df_map_temp)
                        for i, (idx, row) in enumerate(df_map_temp.iterrows()):
                            before = row["before"]
                            after = row["after"]

                            if before and after:
                                mask = df_main["subject_clean"].str.contains(before, regex=False) & \
                                       df_main["subject_clean"].str.contains(after, regex=False)
                            elif before:
                                mask = df_main["subject_clean"].str.contains(before, regex=False)
                            else:
                                if i % max(1, total_patterns // 100) == 0:
                                    step_progress.progress(min(1.0, i / total_patterns))
                                continue

                            new_match = mask & (~matched_flag)

                            for col in selected_output_cols:
                                df_main.loc[new_match, f"mapped_{col}"] = row[col]
                            matched_flag = matched_flag | new_match

                            if i % max(1, total_patterns // 100) == 0:
                                step_progress.progress(min(1.0, i / total_patterns))

                        step_progress.progress(1.0)
                        step_progress.empty()

                        matched_count = matched_flag.sum()
                        unmatched_count = len(df_main) - matched_count
                        status.write(f"✓ Matching complete: {matched_count:,} matched, {unmatched_count:,} no match.")

                        # === Step 3: build output ===
                        status.write("Step 3/3: Finalizing output...")
                        status.update(label=f"✅ {file_name} reconciled successfully!", state="complete")

                        final_out = df_main_original.copy()
                        for col in selected_output_cols:
                            final_out[custom_output_names[col]] = df_main[f"mapped_{col}"]

                        all_results.append({
                            'file_name': file_name,
                            'dataframe': final_out,
                            'row_count': len(final_out)
                        })

                        match_pct = (matched_count / len(df_main) * 100) if len(df_main) > 0 else 0

                        file_status_info.append({
                            "file": file_name,
                            "status": "✅ Success",
                            "rows": len(df_main),
                            "matched": matched_count,
                            "match_pct": f"{match_pct:.1f}%"
                        })

                    except Exception as e:
                        st.error(f"❌ Error processing {file_name}: {str(e)}")
                        file_status_info.append({
                            "file": file_name,
                            "status": "❌ Error",
                            "rows": 0,
                            "reason": str(e)
                        })

                overall_progress_bar.progress(file_idx / len(main_files))

            overall_progress_bar.progress(1.0)
            overall_status_text.write(f"✅ **All {len(main_files)} file(s) processed!**")

            st.session_state.mapping_results = all_results
            st.session_state.file_status = file_status_info

        # --- PERSIST & RENDER RESULTS ---
        if st.session_state.mapping_results:
            all_results = st.session_state.mapping_results
            file_status_info = st.session_state.file_status

            st.success("Batch Processing Complete!")

            st.subheader("Processing Summary")
            summary_df = pd.DataFrame(file_status_info)
            st.dataframe(summary_df, use_container_width=True, hide_index=True)

            # --- INDIVIDUAL FILE PREVIEWS & DOWNLOADS ---
            st.subheader("Preview & Download Results")

            download_files = []

            for result in all_results:
                with st.expander(f"{result['file_name']} - {result['row_count']:,} rows"):
                    st.dataframe(result['dataframe'].head(15), use_container_width=True)

                    csv_buffer = io.StringIO()
                    result['dataframe'].to_csv(csv_buffer, index=False)
                    csv_data = csv_buffer.getvalue()

                    st.download_button(
                        label=f"📥 Download {result['file_name'].split('.')[0]}_mapped.csv",
                        data=csv_data,
                        file_name=f"{result['file_name'].split('.')[0]}_mapped.csv",
                        mime="text/csv",
                        key=f"download_{result['file_name']}"
                    )

                    download_files.append({
                        'file_name': f"{result['file_name'].split('.')[0]}_mapped.csv",
                        'data': csv_data
                    })

            # --- BATCH DOWNLOAD AS ZIP ---
            if len(all_results) > 1:
                st.subheader("Batch Download")

                zip_buffer = io.BytesIO()
                with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
                    for file_info in download_files:
                        zip_file.writestr(file_info['file_name'], file_info['data'])

                zip_buffer.seek(0)

                st.download_button(
                    label="Download All Files as ZIP",
                    data=zip_buffer.getvalue(),
                    file_name=f"mapped_data_batch_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip",
                    mime="application/zip",
                    use_container_width=True
                )

            # --- SUMMARY DASHBOARD ---
            st.subheader("📊 Summary Dashboard")
            st.info("Select columns from your mapped output to generate summary statistics")

            output_columns = list(all_results[0]['dataframe'].columns) if all_results else []

            selected_summary_cols = st.multiselect(
                "Select columns for summary (combination counts)",
                options=output_columns,
                default=output_columns[:min(3, len(output_columns))],
                key="summary_cols_select"
            )

            if selected_summary_cols:
                if st.button("Generate Summary Dashboard", type="secondary"):
                    combo_stats = []

                    for result in all_results:
                        temp_df = result['dataframe'].copy()
                        grouped = temp_df.groupby(selected_summary_cols, dropna=False).size().reset_index(name='Count')
                        grouped['File'] = result['file_name']
                        combo_stats.append(grouped)

                    if combo_stats:
                        summary_dashboard_df = pd.concat(combo_stats, ignore_index=True)

                        st.dataframe(summary_dashboard_df, use_container_width=True, hide_index=True)

                        def create_summary_excel():
                            excel_buffer = io.BytesIO()

                            with pd.ExcelWriter(excel_buffer, engine='openpyxl') as writer:
                                summary_dashboard_df.to_excel(writer, sheet_name='Combination Counts', index=False)

                                worksheet = writer.sheets['Combination Counts']
                                header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                                header_font = Font(bold=True, color="FFFFFF")

                                for cell in worksheet[1]:
                                    cell.fill = header_fill
                                    cell.font = header_font
                                    cell.alignment = Alignment(horizontal="center", vertical="center")

                                for column in worksheet.columns:
                                    max_length = 0
                                    column_letter = column[0].column_letter
                                    for cell in column:
                                        try:
                                            if len(str(cell.value)) > max_length:
                                                max_length = len(str(cell.value))
                                        except:
                                            pass
                                    worksheet.column_dimensions[column_letter].width = max_length + 2

                            excel_buffer.seek(0)
                            return excel_buffer.getvalue()

                        excel_data = create_summary_excel()

                        st.download_button(
                            label="Download Summary Dashboard (Excel)",
                            data=excel_data,
                            file_name=f"summary_dashboard_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True
                        )

            # --- DETAILED SETTINGS USED ---
            st.subheader("Detailed Settings Used")
            st.info("**Column Mapping:**")
            st.write(f"""
            - Subject Column: **{subject_col_name}**
            - Pattern Column: **{pattern_col}**
            - Output Columns: **{', '.join(selected_output_cols)}**
            """)
    else:
        st.warning("Please select at least one output column from your mapping file.")
